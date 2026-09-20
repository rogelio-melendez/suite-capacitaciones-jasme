"""Builds the final branded Excel report from exam config + collected answers.

all_data format: list of rows, each row = [modalidad, nombre, ans1, ans2, ..., ans10]
where ansN is a single lowercase letter ('a'/'b'/'c'/'d') or '' if blank/unrecognized.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

from jasme_style import (
    add_logo, style_title_row, style_header_row, style_clave_row, style_data_row,
    letter_to_full, GRID_BORDER, HEADER_FILL, GREEN_OK, RED_BAD, YELLOW_LIMIT,
    FONT_NAME, BORDER_COLOR,
)

font_arial = Font(name=FONT_NAME)
font_header = Font(name=FONT_NAME, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
thin = Side(style="thin", color=BORDER_COLOR)
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _op_label(config):
    return f"{config['approve_operator']}{config['approve_threshold']}"


def _passes(aciertos, config):
    op, th = config["approve_operator"], config["approve_threshold"]
    if op == ">=":
        return aciertos >= th
    if op == ">":
        return aciertos > th
    if op == "<=":
        return aciertos <= th
    if op == "<":
        return aciertos < th
    return aciertos == th


def build_report_sheet(wb, config, all_data):
    questions, options, answer_key = config["questions"], config["options"], config["answer_key"]
    nq = len(questions)
    ws = wb.create_sheet("Reporte Examen Inicial")
    ncols = 2 + nq

    add_logo(ws)
    style_title_row(ws, 2, ncols, f"Examen Inicial - {config['title']}")

    headers = ["Calificación", "Nombre y Apellido del participante"] + questions
    style_header_row(ws, 3, headers)

    clave_values = ["", "Clave respuestas"] + [letter_to_full(options[i], answer_key[i]) for i in range(nq)]
    style_clave_row(ws, 4, clave_values)

    row = 5
    for entry in all_data:
        name = entry[1]
        answers = entry[2:2 + nq]
        aciertos = sum(1 for a, k in zip(answers, answer_key) if a == k)
        score = round(aciertos * (100 / nq))
        ws.cell(row=row, column=1, value=score).number_format = '0" / 100"'
        ws.cell(row=row, column=2, value=name)
        for j, ans in enumerate(answers):
            ws.cell(row=row, column=3 + j, value=letter_to_full(options[j], ans))
        style_data_row(ws, row, ncols, striped_index=row, center_cols={1})
        row += 1
    last_data_row = row - 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    for j in range(nq):
        ws.column_dimensions[get_column_letter(3 + j)].width = 32
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False

    red_fill = PatternFill(start_color=RED_BAD, end_color=RED_BAD, fill_type="solid")
    for j in range(nq):
        col_letter = get_column_letter(3 + j)
        rng = f"{col_letter}5:{col_letter}{last_data_row}"
        formula = f"{col_letter}5<>{col_letter}$4"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=red_fill))
    return ws


def build_respuestas_sheet(wb, config, all_data):
    questions, answer_key = config["questions"], config["answer_key"]
    nq = len(questions)
    ws = wb.create_sheet("Respuestas")
    n = len(all_data)
    ncols = 2 + nq + 3

    add_logo(ws)
    style_title_row(ws, 2, ncols, f"Respuestas - {config['title']}")

    headers = ["Modalidad", "Nombre del participante"] + [f"P{i}" for i in range(1, nq + 1)] + ["Aciertos", "Total", "Resultado"]
    style_header_row(ws, 3, headers)

    first_data_row = 4
    last_data_row = 3 + n
    key_row = last_data_row + 2

    for i, entry in enumerate(all_data):
        row = first_data_row + i
        modalidad, name = entry[0], entry[1]
        answers = entry[2:2 + nq]
        ws.cell(row=row, column=1, value=modalidad)
        ws.cell(row=row, column=2, value=name)
        for j, ans in enumerate(answers):
            ws.cell(row=row, column=3 + j, value=(ans or "").upper())
        p_cols = [get_column_letter(3 + j) for j in range(nq)]
        formula_parts = [f"IF({p_cols[j]}{row}=$" + f"{p_cols[j]}${key_row},1,0)" for j in range(nq)]
        ws.cell(row=row, column=3 + nq, value="=" + "+".join(formula_parts))
        ws.cell(row=row, column=4 + nq, value=nq)
        op = config["approve_operator"]
        th = config["approve_threshold"]
        col_ac = get_column_letter(3 + nq)
        resultado_formula = f'=IF({col_ac}{row}{op}{th},"Aprobado","No aprobado")'
        ws.cell(row=row, column=5 + nq, value=resultado_formula)
        style_data_row(ws, row, ncols, striped_index=row, center_cols=set(range(1, ncols + 1)) - {2})
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center")

    ws.cell(row=key_row, column=2, value="Respuesta correcta (clave)").font = Font(name=FONT_NAME, bold=True, italic=True)
    for j, k in enumerate(answer_key):
        c = ws.cell(row=key_row, column=3 + j, value=k.upper())
        c.font = Font(name=FONT_NAME, bold=True, italic=True, color="C00000")
        c.alignment = Alignment(horizontal="center")

    green_fill = PatternFill(start_color=GREEN_OK, end_color=GREEN_OK, fill_type="solid")
    red_fill = PatternFill(start_color=RED_BAD, end_color=RED_BAD, fill_type="solid")
    result_col = get_column_letter(5 + nq)
    rng = f"{result_col}{first_data_row}:{result_col}{last_data_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Aprobado"'], fill=green_fill))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No aprobado"'], fill=red_fill))

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 32
    for j in range(nq):
        ws.column_dimensions[get_column_letter(3 + j)].width = 6
    ws.column_dimensions[get_column_letter(3 + nq)].width = 10
    ws.column_dimensions[get_column_letter(4 + nq)].width = 8
    ws.column_dimensions[get_column_letter(5 + nq)].width = 14
    ws.freeze_panes = "C4"
    ws.sheet_view.showGridLines = False
    return ws, first_data_row, last_data_row, nq


def build_resumen_sheet(wb, config, first_data_row, last_data_row, nq):
    ws = wb.create_sheet("Resumen")
    add_logo(ws)
    style_title_row(ws, 2, 5, f"Resumen de resultados - {config['title']}")

    result_col = get_column_letter(5 + nq)
    resp_range = f"Respuestas!${result_col}${first_data_row}:${result_col}${last_data_row}"
    modalidad_range = f"Respuestas!$A${first_data_row}:$A${last_data_row}"

    r = 4
    for i, h in enumerate(["Resultado", "No. de participantes", "% del total"], start=1):
        ws.cell(row=r, column=i, value=h).font = font_header
        ws.cell(row=r, column=i).fill = fill_header
        ws.cell(row=r, column=i).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=i).border = border
    r += 1
    op_label = f"{config['approve_operator']}{config['approve_threshold']}"
    cats = [(f"Aprobado ({op_label})", "Aprobado"), (f"No aprobado", "No aprobado")]
    start_cat_row = r
    for label, key in cats:
        ws.cell(row=r, column=1, value=label).font = font_arial
        ws.cell(row=r, column=2, value=f'=COUNTIF({resp_range},"{key}")').font = font_arial
        ws.cell(row=r, column=3, value=f"=B{r}/$B${start_cat_row + len(cats)}").number_format = "0.0%"
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center" if c > 1 else "left")
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B{start_cat_row}:B{r-1})").font = Font(name=FONT_NAME, bold=True)
    for c in range(1, 4):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center" if c > 1 else "left")
    r += 3

    ws.cell(row=r, column=1, value="Resultados por modalidad").font = Font(name=FONT_NAME, bold=True, size=12)
    r += 1
    hdr2_row = r
    for i, h in enumerate(["Modalidad", "Aprobado", "No aprobado", "Total"], start=1):
        ws.cell(row=r, column=i, value=h).font = font_header
        ws.cell(row=r, column=i).fill = fill_header
        ws.cell(row=r, column=i).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=i).border = border
    r += 1
    modal_data_start = r
    for modal in ["Presencial", "Remoto"]:
        ws.cell(row=r, column=1, value=modal).font = font_arial
        ws.cell(row=r, column=2, value=f'=COUNTIFS({modalidad_range},A{r},{resp_range},"Aprobado")').font = font_arial
        ws.cell(row=r, column=3, value=f'=COUNTIFS({modalidad_range},A{r},{resp_range},"No aprobado")').font = font_arial
        ws.cell(row=r, column=4, value=f"=SUM(B{r}:C{r})").font = font_arial
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center" if c > 1 else "left")
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT_NAME, bold=True)
    for c in range(2, 5):
        col_l = get_column_letter(c)
        ws.cell(row=r, column=c, value=f"=SUM({col_l}{modal_data_start}:{col_l}{modal_data_start+1})").font = Font(name=FONT_NAME, bold=True)
    for c in range(1, 5):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center" if c > 1 else "left")

    ws.column_dimensions["A"].width = 22
    for col_letter in ["B", "C", "D", "E"]:
        ws.column_dimensions[col_letter].width = 15
    ws.sheet_view.showGridLines = False

    pie = PieChart()
    pie.title = "Distribución general de resultados"
    data = Reference(ws, min_col=2, min_row=start_cat_row - 1, max_row=start_cat_row + len(cats) - 1, max_col=2)
    catsref = Reference(ws, min_col=1, min_row=start_cat_row, max_row=start_cat_row + len(cats) - 1)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(catsref)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showCatName = False
    pie.dataLabels.showSerName = False
    pie.dataLabels.showVal = False
    pie.dataLabels.showLegendKey = False
    pie.height = 8
    pie.width = 12
    ws.add_chart(pie, "G4")

    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "Aprobados vs No aprobados por modalidad"
    bar.y_axis.title = "No. de participantes"
    bar.x_axis.title = "Modalidad"
    data = Reference(ws, min_col=2, max_col=3, min_row=hdr2_row, max_row=modal_data_start + 1)
    catsref = Reference(ws, min_col=1, min_row=modal_data_start, max_row=modal_data_start + 1)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(catsref)
    bar.height = 8
    bar.width = 14
    ws.add_chart(bar, "G21")
    return ws


def build_list_sheet(wb, sheet_name, title, config, all_data, want):
    answer_key = config["answer_key"]
    ws = wb.create_sheet(sheet_name)
    add_logo(ws)
    style_title_row(ws, 2, 3, title)

    headers = ["Nombre", "Modalidad", "Aciertos"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h).font = font_header
        ws.cell(row=3, column=i).fill = fill_header
        ws.cell(row=3, column=i).alignment = Alignment(horizontal="center")
        ws.cell(row=3, column=i).border = border

    row = 4
    for entry in all_data:
        modalidad, name = entry[0], entry[1]
        answers = entry[2:2 + len(answer_key)]
        ac = sum(1 for a, k in zip(answers, answer_key) if a == k)
        include = _passes(ac, config) if want == "yes" else not _passes(ac, config)
        if not include:
            continue
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=modalidad)
        ws.cell(row=row, column=3, value=ac)
        style_data_row(ws, row, 3, striped_index=row, center_cols={2, 3})
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
        row += 1

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.sheet_view.showGridLines = False
    return ws


def build_workbook(config, all_data, output_path):
    """config: dict loaded from a config/*.json file.
    all_data: list of [modalidad, nombre, ans1..ansN] (letters, lowercase)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_report_sheet(wb, config, all_data)
    ws_resp, first_row, last_row, nq = build_respuestas_sheet(wb, config, all_data)
    build_resumen_sheet(wb, config, first_row, last_row, nq)
    build_list_sheet(wb, "Aprobados", f"Participantes APROBADOS - {config['title']}", config, all_data, "yes")
    build_list_sheet(wb, "No_Aprobados", f"Participantes NO APROBADOS - {config['title']}", config, all_data, "no")

    wb.save(output_path)
    return output_path
