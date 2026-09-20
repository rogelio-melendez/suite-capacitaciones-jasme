"""JASME branded styling helpers for exam report workbooks (openpyxl)."""
import os
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE_DIR, "assets", "jasme_logo.png")

TITLE_FILL = "0D6596"
HEADER_FILL = "68980D"
CLAVE_FILL = "145A1D"
STRIPE_FILL = "F8F9FA"
BORDER_COLOR = "5B3F86"
GREEN_OK = "C6EFCE"
RED_BAD = "FFC7CE"
YELLOW_LIMIT = "FFEB9C"

FONT_NAME = "Arial"

TITLE_FONT = Font(name=FONT_NAME, bold=True, size=22, color="FFFFFF")
HEADER_FONT_BOLD = Font(name=FONT_NAME, bold=True, size=10, color="FFFFFF")
HEADER_FONT = Font(name=FONT_NAME, size=10, color="FFFFFF")
CLAVE_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=10, color="FFFFFF")
CLAVE_FONT = Font(name=FONT_NAME, size=10, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10, color="000000")
DATA_FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="000000")

_thin_purple = Side(style="thin", color=BORDER_COLOR)
GRID_BORDER = Border(left=_thin_purple, right=_thin_purple, top=_thin_purple, bottom=_thin_purple)


def add_logo(ws, row1_height=117):
    ws.row_dimensions[1].height = row1_height
    if os.path.exists(LOGO_PATH):
        img = XLImage(LOGO_PATH)
        img.width = 210
        img.height = 91
        ws.add_image(img, "A1")


def style_title_row(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = PatternFill(start_color=TITLE_FILL, end_color=TITLE_FILL, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = PatternFill(start_color=TITLE_FILL, end_color=TITLE_FILL, fill_type="solid")


def style_header_row(ws, row, headers, bold_cols=(1, 2), height=42):
    ws.row_dimensions[row].height = height
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=text)
        cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        cell.font = HEADER_FONT_BOLD if i in bold_cols else HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER


def style_clave_row(ws, row, values, height=45):
    ws.row_dimensions[row].height = height
    for i, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=val if val else None)
        cell.fill = PatternFill(start_color=CLAVE_FILL, end_color=CLAVE_FILL, fill_type="solid")
        cell.font = CLAVE_LABEL_FONT if i == 2 else CLAVE_FONT
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = GRID_BORDER


def style_data_row(ws, row, ncols, striped_index, center_cols=None, bold_col1=False):
    fill_color = STRIPE_FILL if striped_index % 2 == 0 else "FFFFFF"
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    center_cols = center_cols or set()
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = DATA_FONT_BOLD if (c == 1 and bold_col1) else DATA_FONT
        cell.alignment = Alignment(horizontal="center" if c in center_cols else "left",
                                    vertical="center", wrap_text=True)
        cell.border = GRID_BORDER


def letter_to_full(options_dict, letter):
    letter = (letter or "").strip().lower()
    text = options_dict.get(letter, "")
    return f"{letter}) {text}" if text else (letter.upper() if letter else "")
