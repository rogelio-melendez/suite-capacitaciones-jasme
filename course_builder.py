# -*- coding: utf-8 -*-
"""Generalized JASME course-deck builder.

Takes a structured course_data dict (temario, objetivo, per-topic content +
dynamics, exam questions) and assembles it into the JASME template, using
the same technique proven on the NOM-010 course: duplicate a handful of
"role" slides from the template as many times as needed, then fill in text
via python-pptx.

This module knows exactly one thing about the template's internals: which
slideN.xml plays which role. That map is defined once, in TEMPLATE_MAP,
against the bundled assets/Plantilla_JASME.pptx.
"""
import copy
import os
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

import slide_ops

NSMAP = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}

# Roles, expressed as the slideN.xml filename that plays them in the
# bundled JASME template (Plantilla_JASME.pptx). If the template file is
# ever replaced, these must be re-verified against the new file's numbering.
ROLE_TITLE_IMG = "slide19.xml"      # numbered section title w/ image (ctrTitle + pic)
ROLE_CONTENT = "slide24.xml"        # title + body(idx1) + label(idx14)
ROLE_DINAMICA = "slide26.xml"       # "Dinámica" + instruction
ROLE_EJERCICIO_Q = "slide27.xml"    # "Ejercicio N" + statement only
ROLE_EJERCICIO_A = "slide28.xml"    # "Ejercicio N" + statement + verdict/explanation
ROLE_EXAM_PAGE = "slide17.xml"      # exam page: idx14 body w/ buAutoNum options

# Slides removed wholesale and rebuilt: original demo topics (19-74), and the
# two fixed-size exam blocks (17-18 initial, 80-83 final).
SLIDES_TO_REMOVE = (
    [f"slide{n}.xml" for n in range(17, 75)] +
    [f"slide{n}.xml" for n in range(80, 84)]
)

# Anchors in the ORIGINAL template's <p:sldIdLst> we insert new content after.
ANCHOR_AFTER_EXAM_INICIAL_INTRO = "slide16.xml"   # exam-inicial pages go here
ANCHOR_AFTER_EXAM_FINAL_INTRO = "slide79.xml"     # exam-final pages go here
# Topic blocks are inserted right after the last exam-inicial page we create.


def _ph(slide, idx):
    for shp in slide.placeholders:
        if shp.placeholder_format.idx == idx:
            return shp
    raise KeyError(f"No placeholder idx={idx}")


def _set_simple_text(shape, text, size=None):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if not p.runs:
        run = p.add_run()
    else:
        run = p.runs[0]
        for extra in p.runs[1:]:
            extra._r.getparent().remove(extra._r)
    run.text = text
    if size is not None:
        run.font.size = Pt(size)


def _set_multiline(shape, lines, size=20):
    tf = shape.text_frame
    txBody = tf._txBody
    paras = txBody.findall('a:p', NSMAP)
    template_p = paras[0]
    template_pPr = template_p.find('a:pPr', NSMAP)
    for p_el in paras[1:]:
        txBody.remove(p_el)
    for child in list(template_p):
        template_p.remove(child)
    if template_pPr is not None:
        template_p.append(copy.deepcopy(template_pPr))

    def make_run(parent, text):
        r = etree.SubElement(parent, '{%s}r' % NSMAP['a'])
        rPr = etree.SubElement(r, '{%s}rPr' % NSMAP['a'])
        rPr.set('lang', 'es-MX'); rPr.set('sz', str(int(size * 100))); rPr.set('dirty', '0')
        t = etree.SubElement(r, '{%s}t' % NSMAP['a']); t.text = text

    make_run(template_p, lines[0] if lines else "")
    for line in lines[1:]:
        p_el = etree.SubElement(txBody, '{%s}p' % NSMAP['a'])
        if template_pPr is not None:
            p_el.append(copy.deepcopy(template_pPr))
        make_run(p_el, line)


def _set_two_paragraphs(shape, verdict, explanation, size=24):
    tf = shape.text_frame
    txBody = tf._txBody
    paras = txBody.findall('a:p', NSMAP)
    p1 = paras[0]
    for child in list(p1):
        p1.remove(child)
    r1 = etree.SubElement(p1, '{%s}r' % NSMAP['a'])
    rPr1 = etree.SubElement(r1, '{%s}rPr' % NSMAP['a'])
    rPr1.set('lang', 'es-MX'); rPr1.set('sz', str(size*100)); rPr1.set('b', '1'); rPr1.set('u', 'sng'); rPr1.set('dirty', '0')
    etree.SubElement(r1, '{%s}t' % NSMAP['a']).text = verdict

    if len(paras) > 1:
        p2 = paras[1]
        for extra in paras[2:]:
            txBody.remove(extra)
    else:
        p2 = copy.deepcopy(p1)
        txBody.append(p2)
    for child in list(p2):
        p2.remove(child)
    r2 = etree.SubElement(p2, '{%s}r' % NSMAP['a'])
    rPr2 = etree.SubElement(r2, '{%s}rPr' % NSMAP['a'])
    rPr2.set('lang', 'es-MX'); rPr2.set('sz', str(size*100)); rPr2.set('dirty', '0')
    etree.SubElement(r2, '{%s}t' % NSMAP['a']).text = explanation


def _find_content_shape(slide):
    """The idx=1 body placeholder (never idx=11 footer / idx=14 label)."""
    return _ph(slide, 1)


def _set_exam_page(slide, preamble_lines, questions, start_num, size=2000):
    """questions: list of (question_text, [opt_a, opt_b, opt_c, opt_d])"""
    shape = _ph(slide, 14)
    txBody = shape.text_frame._txBody
    for p_el in txBody.findall('a:p', NSMAP):
        txBody.remove(p_el)

    def add(xml_str):
        wrapped = xml_str.replace(
            '<a:p>', '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">', 1
        )
        txBody.append(etree.fromstring(wrapped.encode('utf-8')))

    for line in preamble_lines:
        add(f'<a:p><a:pPr algn="l"/><a:r><a:rPr lang="es-MX" sz="1800" dirty="0"/><a:t>{line}</a:t></a:r></a:p>')
    for i, (qtext, opts) in enumerate(questions):
        num = start_num + i
        add(f'<a:p><a:pPr algn="l"/><a:r><a:rPr lang="es-MX" sz="{size}" b="1" dirty="0"/>'
            f'<a:t>{num}.- {qtext}</a:t></a:r><a:endParaRPr lang="es-MX" sz="{size}" dirty="0"/></a:p>')
        for opt in opts:
            add('<a:p><a:pPr marL="457200" indent="-457200" algn="l"><a:buAutoNum type="alphaLcParenR"/></a:pPr>'
                f'<a:r><a:rPr lang="es-MX" sz="{size}" dirty="0"/><a:t>{opt}</a:t></a:r></a:p>')
        add(f'<a:p><a:pPr algn="l"/><a:endParaRPr lang="es-MX" sz="{size}" dirty="0"/></a:p>')


def _chunk_questions(questions, per_page=5):
    return [questions[i:i + per_page] for i in range(0, len(questions), per_page)]


def build_course(template_path, output_path, course_data, progress_callback=None):
    """course_data schema -- see course_ai.py for how each field is produced:
    {
      "title": "NORMA Oficial Mexicana ...",
      "temario": ["Tema 1", "Tema 2", ...],
      "objetivo": "Al finalizar el curso ...",
      "topics": [
        {
          "title": "1. Tema 1",
          "content_slides": [["bullet1", "bullet2"], ["bullet1", ...], ...],
          "dynamic": {
             "instruction": "...",
             "exercises": [{"statement":.., "verdict":.., "explanation":..}, ...]  # up to 4
          } | None
        }, ...
      ],
      "exam_inicial": [("pregunta", ["opt_a","opt_b","opt_c","opt_d"]), ...],
      "exam_final": [("pregunta", [...]), ...],
    }
    """
    def report(msg):
        if progress_callback:
            progress_callback(msg)

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        report("Descomprimiendo plantilla...")
        unpacked = slide_ops.extract_pptx(template_path, tmp / "unpacked")

        report("Quitando los temas de ejemplo de la plantilla...")
        slide_ops.remove_slides(unpacked, SLIDES_TO_REMOVE)

        # ---- Build exam-inicial pages ----
        report("Generando páginas del examen inicial...")
        cursor = ANCHOR_AFTER_EXAM_INICIAL_INTRO
        inicial_pages = []
        for page_qs in _chunk_questions(course_data["exam_inicial"]):
            new_file = slide_ops.duplicate_slide(unpacked, ROLE_EXAM_PAGE, after=cursor)
            inicial_pages.append(new_file)
            cursor = new_file

        # ---- Build topic blocks ----
        topic_slide_map = []  # list of dicts describing what to fill later
        for topic in course_data["topics"]:
            report(f"Construyendo tema: {topic['title']}...")
            title_file = slide_ops.duplicate_slide(unpacked, ROLE_TITLE_IMG, after=cursor)
            cursor = title_file
            content_files = []
            for _ in topic["content_slides"]:
                cf = slide_ops.duplicate_slide(unpacked, ROLE_CONTENT, after=cursor)
                content_files.append(cf)
                cursor = cf
            dyn_file = None
            exercise_files = []
            if topic.get("dynamic"):
                dyn_file = slide_ops.duplicate_slide(unpacked, ROLE_DINAMICA, after=cursor)
                cursor = dyn_file
                for _ in topic["dynamic"]["exercises"]:
                    qf = slide_ops.duplicate_slide(unpacked, ROLE_EJERCICIO_Q, after=cursor)
                    cursor = qf
                    af = slide_ops.duplicate_slide(unpacked, ROLE_EJERCICIO_A, after=cursor)
                    cursor = af
                    exercise_files.append((qf, af))
            topic_slide_map.append({
                "topic": topic, "title_file": title_file, "content_files": content_files,
                "dyn_file": dyn_file, "exercise_files": exercise_files,
            })

        # ---- Build exam-final pages ----
        report("Generando páginas del examen final...")
        final_cursor = ANCHOR_AFTER_EXAM_FINAL_INTRO
        final_pages = []
        for page_qs in _chunk_questions(course_data["exam_final"]):
            new_file = slide_ops.duplicate_slide(unpacked, ROLE_EXAM_PAGE, after=final_cursor)
            final_pages.append(new_file)
            final_cursor = new_file

        report("Limpiando archivos huérfanos...")
        slide_ops.prune_unreferenced_slides(unpacked)

        # Ground truth for slide order -- captured BEFORE repacking, since
        # python-pptx's own part.partname renumbers on load and can't be
        # trusted to recover which physical slideN.xml is which.
        ordered_files = slide_ops.get_ordered_slide_files(unpacked)

        report("Empacando y abriendo con python-pptx para llenar el texto...")
        tmp_pptx = tmp / "assembled.pptx"
        slide_ops.repack_pptx(unpacked, tmp_pptx)

        # =========================================================
        # Content fill pass (python-pptx)
        # =========================================================
        prs = Presentation(str(tmp_pptx))
        if len(prs.slides) != len(ordered_files):
            raise RuntimeError(
                f"Slide count mismatch after assembly: expected {len(ordered_files)}, "
                f"python-pptx sees {len(prs.slides)}. The deck may be corrupt."
            )
        slides_by_file = dict(zip(ordered_files, prs.slides))

        def S(fname):
            return slides_by_file[fname]

        report("Escribiendo título, temario y objetivo...")
        # Title slide (slide1.xml)
        _set_simple_text(_ph(S("slide1.xml"), 1), course_data["title"], size=36)
        # Bienvenida (slide10.xml) -- append course title to the greeting line
        for shp in S("slide10.xml").shapes:
            if shp.has_text_frame and "Sean bienvenidos" in shp.text_frame.text:
                shp.text_frame.paragraphs[0].runs[0].text = f"Sean bienvenidos al curso de {course_data['title']}"
        # Temario (slide12.xml) -- content lives in idx=14, NOT idx=11 (that's the footer)
        _set_multiline(_ph(S("slide12.xml"), 14), course_data["temario"], size=20)
        # Objetivo del curso (slide13.xml)
        _set_simple_text(_find_content_shape(S("slide13.xml")), course_data["objetivo"], size=22)

        report("Escribiendo el examen inicial...")
        for page_idx, page_file in enumerate(inicial_pages):
            page_qs = _chunk_questions(course_data["exam_inicial"])[page_idx]
            start_num = page_idx * 5 + 1
            preamble = (["Nombre del participante_______________ ", "Fecha: ___________",
                         "", "Seleccione la respuesta correcta."] if page_idx == 0 else [])
            _set_exam_page(S(page_file), preamble, page_qs, start_num)

        report("Escribiendo cada tema, su contenido y su dinámica...")
        for entry in topic_slide_map:
            topic = entry["topic"]
            _set_simple_text(S(entry["title_file"]).shapes.title, topic["title"])
            for cf, bullets in zip(entry["content_files"], topic["content_slides"]):
                sl = S(cf)
                # clear the stray "Definición 1:" label inherited from the role slide
                try:
                    _set_simple_text(_ph(sl, 14), "")
                except KeyError:
                    pass
                if topic.get("content_titles"):
                    idx = entry["content_files"].index(cf)
                    _set_simple_text(sl.shapes.title, topic["content_titles"][idx])
                _set_multiline(_find_content_shape(sl), bullets, size=20)
            if entry["dyn_file"]:
                _set_simple_text(_ph(S(entry["dyn_file"]), 1), topic["dynamic"]["instruction"], size=28)
                for i, ((qf, af), ex) in enumerate(zip(entry["exercise_files"], topic["dynamic"]["exercises"]), start=1):
                    _set_simple_text(S(qf).shapes.title, f"Ejercicio {i}")
                    _set_simple_text(_ph(S(qf), 14), ex["statement"], size=24)
                    _set_simple_text(S(af).shapes.title, f"Ejercicio {i}")
                    _set_simple_text(_ph(S(af), 14), ex["statement"], size=24)
                    _set_two_paragraphs(_ph(S(af), 1), ex["verdict"], ex["explanation"], size=24)

        report("Escribiendo el examen final...")
        for page_idx, page_file in enumerate(final_pages):
            page_qs = _chunk_questions(course_data["exam_final"])[page_idx]
            start_num = page_idx * 5 + 1
            preamble = (["Nombre del participante_______________ ", "Fecha: ___________",
                         "", "Seleccione la respuesta correcta."] if page_idx == 0 else [])
            _set_exam_page(S(page_file), preamble, page_qs, start_num)

        report("Guardando el archivo final...")
        prs.save(str(output_path))

    return output_path
