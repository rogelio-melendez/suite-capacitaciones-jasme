# -*- coding: utf-8 -*-
"""Renders a full course deck from nothing but structured content + a
DeckStyle (colors/fonts). This is the render path that can never break from
an external template changing, because there is no external template --
every shape is placed by this code.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)


def _rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # layout 6 = blank, on the default template


def _set_background(slide, color_hex):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(color_hex)


def _add_footer(slide, style, page_num, course_title, light=False):
    color = style.text_light if light else style.text_dark
    tb = slide.shapes.add_textbox(MARGIN, SLIDE_H - Inches(0.4), SLIDE_W - 2 * MARGIN, Inches(0.3))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"{course_title}"
    run.font.size = Pt(9)
    run.font.name = style.body_font
    run.font.color.rgb = _rgb(color)
    p.alignment = PP_ALIGN.LEFT

    tb2 = slide.shapes.add_textbox(SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.4), Inches(0.8), Inches(0.3))
    p2 = tb2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = str(page_num)
    run2.font.size = Pt(9)
    run2.font.name = style.body_font
    run2.font.color.rgb = _rgb(color)
    p2.alignment = PP_ALIGN.RIGHT


def _textbox(slide, left, top, width, height, text, size, color, bold=False, font=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, underline=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.name = font
    run.font.color.rgb = _rgb(color)
    return tb


def _bullets(slide, left, top, width, height, lines, size, color, font, line_spacing=1.25):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {line}"
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = _rgb(color)
    return tb


def build_title_slide(prs, style, course_title):
    slide = _blank_slide(prs)
    _set_background(slide, style.primary)
    _textbox(slide, Inches(1), Inches(2.7), Inches(11.3), Inches(1.2),
              "CURSO", 24, style.text_light, bold=True, font=style.heading_font, align=PP_ALIGN.CENTER)
    _textbox(slide, Inches(1), Inches(3.3), Inches(11.3), Inches(1.8),
              course_title, 40, style.text_light, bold=True, font=style.heading_font, align=PP_ALIGN.CENTER)

def build_temario_slide(prs, style, temario, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.9),
              "Tabla de contenido del curso", 32, style.primary, bold=True, font=style.heading_font)
    lines = [f"{i + 1}.  {t}" for i, t in enumerate(temario)]
    _bullets(slide, MARGIN, Inches(1.7), SLIDE_W - 2 * MARGIN, Inches(5),
             lines, 18, style.text_dark, style.body_font, line_spacing=1.4)
    _add_footer(slide, style, page_num, course_title)


def build_objetivo_slide(prs, style, objetivo, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.9),
              "Objetivo de aprendizaje", 32, style.primary, bold=True, font=style.heading_font)
    _textbox(slide, MARGIN, Inches(1.8), SLIDE_W - 2 * MARGIN, Inches(4.5),
              objetivo, 20, style.text_dark, font=style.body_font)
    _add_footer(slide, style, page_num, course_title)


def build_topic_title_slide(prs, style, topic_title, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.primary)
    _textbox(slide, Inches(1), Inches(3.0), Inches(11.3), Inches(1.5),
              topic_title, 34, style.text_light, bold=True, font=style.heading_font, align=PP_ALIGN.CENTER)
    _add_footer(slide, style, page_num, course_title, light=True)


def build_content_slide(prs, style, subtitle, bullets, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.9),
              subtitle, 28, style.primary, bold=True, font=style.heading_font)
    _bullets(slide, MARGIN, Inches(1.6), SLIDE_W - 2 * MARGIN, Inches(5.2),
             bullets, 18, style.text_dark, style.body_font)
    _add_footer(slide, style, page_num, course_title)


def build_dinamica_slide(prs, style, instruction, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.accent)
    _textbox(slide, Inches(1.5), Inches(2.3), Inches(10.3), Inches(1.0),
              "Dinámica", 36, style.text_light, bold=True, font=style.heading_font, align=PP_ALIGN.CENTER)
    _textbox(slide, Inches(2), Inches(3.3), Inches(9.3), Inches(2),
              instruction, 20, style.text_light, font=style.body_font, align=PP_ALIGN.CENTER)
    _add_footer(slide, style, page_num, course_title, light=True)


def build_ejercicio_q_slide(prs, style, num, statement, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.9),
              f"Ejercicio {num}", 28, style.primary, bold=True, font=style.heading_font)
    _textbox(slide, MARGIN, Inches(2.2), SLIDE_W - 2 * MARGIN, Inches(2.5),
              statement, 22, style.text_dark, font=style.body_font)
    _add_footer(slide, style, page_num, course_title)


def build_ejercicio_a_slide(prs, style, num, statement, verdict, explanation, course_title, page_num):
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.9),
              f"Ejercicio {num}", 28, style.primary, bold=True, font=style.heading_font)
    _textbox(slide, MARGIN, Inches(1.5), SLIDE_W - 2 * MARGIN, Inches(1.5),
              statement, 20, style.text_dark, font=style.body_font)
    _textbox(slide, MARGIN, Inches(3.1), SLIDE_W - 2 * MARGIN, Inches(0.6),
              verdict, 22, style.accent, bold=True, underline=True, font=style.heading_font)
    _textbox(slide, MARGIN, Inches(3.8), SLIDE_W - 2 * MARGIN, Inches(2.5),
              explanation, 18, style.text_dark, font=style.body_font)
    _add_footer(slide, style, page_num, course_title)


def build_exam_slide(prs, style, title, preamble_lines, questions, start_num, course_title, page_num):
    """questions: list of (question_text, [opt_a, opt_b, opt_c, opt_d])"""
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.4), SLIDE_W - 2 * MARGIN, Inches(0.7),
              title, 28, style.primary, bold=True, font=style.heading_font, align=PP_ALIGN.CENTER)

    col_width = (SLIDE_W - 2 * MARGIN - Inches(0.4)) / 2
    left_box = slide.shapes.add_textbox(MARGIN, Inches(1.2), col_width, Inches(5.8))
    right_box = slide.shapes.add_textbox(MARGIN + col_width + Inches(0.4), Inches(1.2), col_width, Inches(5.8))
    left_tf, right_tf = left_box.text_frame, right_box.text_frame
    left_tf.word_wrap = True
    right_tf.word_wrap = True

    def write_preamble(tf):
        for i, line in enumerate(preamble_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size = Pt(13)
            run.font.name = style.body_font
            run.font.color.rgb = _rgb(style.text_dark)

    def write_question(tf, first, num, qtext, opts):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = f"{num}.- {qtext}"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.name = style.body_font
        run.font.color.rgb = _rgb(style.text_dark)
        for j, opt in enumerate(opts):
            letter = "abcd"[j]
            po = tf.add_paragraph()
            po.level = 1
            run_o = po.add_run()
            run_o.text = f"{letter}) {opt}"
            run_o.font.size = Pt(13)
            run_o.font.name = style.body_font
            run_o.font.color.rgb = _rgb(style.text_dark)

    half = (len(questions) + 1) // 2
    left_qs, right_qs = questions[:half], questions[half:]

    first_left = True
    if preamble_lines:
        write_preamble(left_tf)
        first_left = False
    for i, (qtext, opts) in enumerate(left_qs):
        write_question(left_tf, first_left and i == 0, start_num + i, qtext, opts)

    first_right = True
    for i, (qtext, opts) in enumerate(right_qs):
        write_question(right_tf, first_right and i == 0, start_num + half + i, qtext, opts)
        first_right = False

    _add_footer(slide, style, page_num, course_title)


def build_generic_text_slide(prs, style, title, body_text, course_title, page_num):
    """Simple fallback slide for generic content (rules, welcome, etc.) when
    the user writes it as plain text instead of uploading a library block."""
    slide = _blank_slide(prs)
    _set_background(slide, style.background)
    _textbox(slide, MARGIN, Inches(0.6), SLIDE_W - 2 * MARGIN, Inches(1.0),
              title, 30, style.primary, bold=True, font=style.heading_font)
    _textbox(slide, MARGIN, Inches(1.8), SLIDE_W - 2 * MARGIN, Inches(5),
              body_text, 18, style.text_dark, font=style.body_font)
    _add_footer(slide, style, page_num, course_title)


def build_course_freeform(output_path, course_data, style):
    """course_data: same schema as course_builder.build_course's, minus any
    template dependency."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    course_title = course_data["title"]
    page = 1

    build_title_slide(prs, style, course_title); page += 1
    build_temario_slide(prs, style, course_data["temario"], course_title, page); page += 1
    build_objetivo_slide(prs, style, course_data["objetivo"], course_title, page); page += 1

    if course_data.get("exam_inicial"):
        questions = course_data["exam_inicial"]
        for i in range(0, len(questions), 5):
            chunk = questions[i:i + 5]
            preamble = (["Nombre del participante: ______________________", "Fecha: ____________",
                         "Seleccione la respuesta correcta."] if i == 0 else [])
            build_exam_slide(prs, style, "Examen inicial", preamble, chunk, i + 1, course_title, page)
            page += 1

    for topic in course_data["topics"]:
        build_topic_title_slide(prs, style, topic["title"], course_title, page); page += 1
        for subtitle, bullets in zip(topic["content_titles"], topic["content_slides"]):
            build_content_slide(prs, style, subtitle, bullets, course_title, page); page += 1
        dyn = topic.get("dynamic")
        if dyn:
            build_dinamica_slide(prs, style, dyn["instruction"], course_title, page); page += 1
            for i, ex in enumerate(dyn["exercises"], start=1):
                build_ejercicio_q_slide(prs, style, i, ex["statement"], course_title, page); page += 1
                build_ejercicio_a_slide(prs, style, i, ex["statement"], ex["verdict"], ex["explanation"],
                                         course_title, page); page += 1

    if course_data.get("exam_final"):
        questions = course_data["exam_final"]
        for i in range(0, len(questions), 5):
            chunk = questions[i:i + 5]
            preamble = (["Nombre del participante: ______________________", "Fecha: ____________",
                         "Seleccione la respuesta correcta."] if i == 0 else [])
            build_exam_slide(prs, style, "Examen final", preamble, chunk, i + 1, course_title, page)
            page += 1

    prs.save(output_path)
    return output_path
