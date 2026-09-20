"""Extracts plain text from a source document (PDF, DOCX, or TXT) so it can
be handed to the AI planning steps as the raw material for a course."""
import io


def extract_text(uploaded_file, filename):
    """uploaded_file: a file-like object (e.g. Streamlit's UploadedFile).
    filename: original filename, used to pick the right extractor."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    data = uploaded_file.read()

    if ext == "pdf":
        return _extract_pdf(data)
    elif ext in ("docx",):
        return _extract_docx(data)
    elif ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Formato no soportado: .{ext}. Usa PDF, DOCX o TXT.")


def _extract_pdf(data: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n\n".join(text_parts)


def _extract_docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
